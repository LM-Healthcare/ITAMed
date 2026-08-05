"""
==============================================================================
ITAMed — Apply Question Type Annotation to All Dataset Files
==============================================================================

Purpose:
    Reads the question type classification (case-based / knowledge-based) from
    results/question_type/question_types.json and applies it to all official
    dataset files (IT + EN, JSON + XLSX, per-year + complete).

Input:
    - results/question_type/question_types.json
    - Dataset/IT/json/ and Dataset/EN/json/ (JSON files)
    - Dataset/IT/xlsx/ and Dataset/EN/xlsx/ (XLSX files)

Output:
    Updated dataset files with the "question_type" field added.

Usage:
    python scripts/apply_question_type.py

Author: ITAMed Dataset Team
==============================================================================
"""

import os
import json

import openpyxl

# ==============================================================================
# Configuration
# ==============================================================================

QC_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPO = os.path.dirname(QC_DIR)

IT_JSON_DIR = os.path.join(REPO, "Dataset", "IT", "json")
IT_XLSX_DIR = os.path.join(REPO, "Dataset", "IT", "xlsx")
EN_JSON_DIR = os.path.join(REPO, "Dataset", "EN", "json")
EN_XLSX_DIR = os.path.join(REPO, "Dataset", "EN", "xlsx")
TYPES_FILE = os.path.join(QC_DIR, "results", "question_type", "question_types.json")

YEARS = list(range(2017, 2026))

# Translation for EN files
TYPE_IT_TO_EN = {
    "case-based": "case-based",
    "knowledge-based": "knowledge-based",
}


# ==============================================================================
# JSON Update
# ==============================================================================


def apply_to_json(json_dir, suffix, types):
    """Add question_type field to JSON files."""
    updated = 0

    for year in YEARS:
        path = os.path.join(json_dir, f"ITAMed_{year}{suffix}.json")
        if not os.path.exists(path):
            continue
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        changed = False
        for item in data:
            code = item["question_code"]
            q_type = types.get(code, "")
            if q_type and item.get("question_type") != q_type:
                item["question_type"] = q_type
                changed = True
                updated += 1
        if changed:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)

    # Complete file
    path = os.path.join(json_dir, f"ITAMed_complete{suffix}.json")
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        changed = False
        for item in data:
            code = item["question_code"]
            q_type = types.get(code, "")
            if q_type and item.get("question_type") != q_type:
                item["question_type"] = q_type
                changed = True
                updated += 1
        if changed:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)

    return updated


# ==============================================================================
# XLSX Update
# ==============================================================================


def apply_to_xlsx(xlsx_dir, suffix, types):
    """Add question_type column (L) to XLSX files."""
    updated = 0
    TYPE_COL = 12  # Column L (0-indexed = 11, 1-indexed = 12)

    for name in [f"ITAMed_{y}{suffix}.xlsx" for y in YEARS] + [f"ITAMed_complete{suffix}.xlsx"]:
        path = os.path.join(xlsx_dir, name)
        if not os.path.exists(path):
            continue

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        changed = False

        # Add header if not present
        header_cell = ws.cell(row=1, column=TYPE_COL)
        if header_cell.value != "question_type":
            header_cell.value = "question_type"
            changed = True

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            code = str(row[2].value) if row[2].value else ""
            q_type = types.get(code, "")
            cell = ws.cell(row=row[0].row, column=TYPE_COL)
            if q_type and cell.value != q_type:
                cell.value = q_type
                changed = True
                updated += 1

        if changed:
            wb.save(path)

    return updated


# ==============================================================================
# Main
# ==============================================================================


def main():
    print("=" * 70)
    print("ITAMed — Apply Question Type Annotation")
    print("=" * 70)

    # Load classification
    with open(TYPES_FILE, "r", encoding="utf-8") as f:
        types = json.load(f)
    print(f"\nLoaded {len(types)} question type annotations")

    # Apply to IT JSON
    print("\nApplying to IT JSON files...")
    n_it_json = apply_to_json(IT_JSON_DIR, "", types)
    print(f"  Updates: {n_it_json}")

    # Apply to EN JSON
    print("Applying to EN JSON files...")
    n_en_json = apply_to_json(EN_JSON_DIR, "_EN", types)
    print(f"  Updates: {n_en_json}")

    # Apply to IT XLSX
    print("Applying to IT XLSX files...")
    n_it_xlsx = apply_to_xlsx(IT_XLSX_DIR, "", types)
    print(f"  Updates: {n_it_xlsx}")

    # Apply to EN XLSX
    print("Applying to EN XLSX files...")
    n_en_xlsx = apply_to_xlsx(EN_XLSX_DIR, "_EN", types)
    print(f"  Updates: {n_en_xlsx}")

    total = n_it_json + n_en_json + n_it_xlsx + n_en_xlsx
    print(f"\n{'='*70}")
    print(f"DONE — Total field updates: {total}")
    print(f"{'='*70}")


if __name__ == "__main__":
    main()
