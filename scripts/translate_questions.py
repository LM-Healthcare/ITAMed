"""
==============================================================================
ITAMed — Medical Question Translation Script (Italian → English)
==============================================================================

Purpose:
    Translates all classified SSM medical exam questions from Italian to
    English using the Anthropic Claude API.

    The translation is performed by an LLM configured as an expert medical
    translator, ensuring:
    - Correct English medical terminology (not lay terms)
    - Native-level fluency (not literal translation)
    - Preservation of all clinical details (lab values, dosages, anatomy)
    - USMLE-style English conventions

Input:
    Classified XLSX files in 'Data/Classified/' containing Italian questions
    with specialty annotations.
    Expected format: columns [Codice Domanda, Domanda, Risposta A–E,
                              Categoria, Immagine, Categoria Immagine]

Output:
    Translated XLSX files in 'Data/Translated_EN/' with English text
    and translated metadata (category names, image flags).
    Intermediate JSON files are saved for resume support.

Requirements:
    - anthropic>=0.42 (Claude API client)
    - openpyxl (XLSX read/write)
    - ANTHROPIC_API_KEY environment variable must be set

Usage:
    export ANTHROPIC_API_KEY="sk-ant-..."
    python translate_questions.py

Author: ITAMed Dataset Team
==============================================================================
"""

import os
import sys
import json
import time
import re

import openpyxl
import anthropic

# ==============================================================================
# Configuration
# ==============================================================================

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CLASSIFIED_DIR = os.path.join(BASE_DIR, "Data", "Classified")
OUTPUT_DIR = os.path.join(BASE_DIR, "Data", "Translated_EN")
INTERMEDIATE_DIR = os.path.join(OUTPUT_DIR, "intermediate")

MODEL = "claude-opus-4-8"
BATCH_SIZE = 5
YEARS = list(range(2017, 2026))
EXPECTED_QUESTIONS = 140

# ==============================================================================
# Translation Prompt
# ==============================================================================

SYSTEM_PROMPT = """You are an expert medical translator with deep specialization in Italian-to-English translation of clinical and biomedical texts. You hold equivalent qualifications to a certified medical translator (CMT) and have extensive experience translating medical examination questions, clinical case descriptions, and pharmacological/pathological terminology.

Your task is to translate Italian medical exam questions into fluent, native-level English. Your translations must:

1. Be medically accurate: Use the correct English medical terminology (e.g., "infarto miocardico" → "myocardial infarction", not "heart attack"; "emocromo" → "complete blood count", not "blood test").

2. Sound native: The result should read as if it were originally written in English by a native-speaking medical professional. Avoid calques, literal translations, or awkward phrasing.

3. Preserve clinical precision: Do not simplify, paraphrase, or omit any clinical detail. Every lab value, dosage, anatomical reference, and temporal detail must be faithfully preserved.

4. Maintain the question structure: Keep the same format — a clinical vignette followed by a question stem, with five answer options (A through E). Do not reorder or modify the answer options.

5. Handle idiomatic Italian medical expressions correctly: Translate them into their standard English equivalents as used in international medical literature and examinations (e.g., USMLE-style English).

6. Preserve formatting: If the original has numbered items, bullet points, or specific formatting markers, keep them.

Respond ONLY with a valid JSON array containing the translated questions. Each element must have:
- "question": the translated question text (clinical vignette + question stem)
- "A": translated answer option A
- "B": translated answer option B
- "C": translated answer option C
- "D": translated answer option D
- "E": translated answer option E

Do NOT include any commentary, explanation, or text outside the JSON array."""


# ==============================================================================
# Translation Logic
# ==============================================================================


def translate_batch(client: anthropic.Anthropic, questions_batch: list, batch_start_idx: int) -> list | None:
    """
    Translate a batch of questions using the Claude API.

    Args:
        client: Anthropic API client instance.
        questions_batch: List of question dicts with Italian text.
        batch_start_idx: Starting index of this batch (for logging).

    Returns:
        List of translated question dicts, or None if all retries failed.
    """
    
    # Build the user prompt
    prompt_parts = []
    for i, q in enumerate(questions_batch):
        prompt_parts.append(f"--- Question {batch_start_idx + i + 1} ---")
        prompt_parts.append(f"Domanda: {q['domanda']}")
        prompt_parts.append(f"A) {q['A']}")
        prompt_parts.append(f"B) {q['B']}")
        prompt_parts.append(f"C) {q['C']}")
        prompt_parts.append(f"D) {q['D']}")
        prompt_parts.append(f"E) {q['E']}")
        prompt_parts.append("")
    
    user_prompt = "Translate the following Italian medical exam questions into English:\n\n" + "\n".join(prompt_parts)
    
    max_retries = 3
    for attempt in range(max_retries):
        try:
            response = client.messages.create(
                model="claude-opus-4-8",
                max_tokens=4096,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": user_prompt}]
            )
            
            response_text = response.content[0].text.strip()
            
            # Parse JSON response
            # Try to find JSON array in the response
            json_match = re.search(r'\[.*\]', response_text, re.DOTALL)
            if json_match:
                results = json.loads(json_match.group())
            else:
                results = json.loads(response_text)
            
            if len(results) == len(questions_batch):
                return results
            else:
                print(f"    WARNING: Expected {len(questions_batch)} results, got {len(results)}. Retrying...")
                
        except json.JSONDecodeError as e:
            print(f"    JSON parse error (attempt {attempt+1}): {e}")
            if attempt < max_retries - 1:
                time.sleep(2)
        except anthropic.APIError as e:
            print(f"    API error (attempt {attempt+1}): {e}")
            if attempt < max_retries - 1:
                time.sleep(5)
        except Exception as e:
            print(f"    Unexpected error (attempt {attempt+1}): {e}")
            if attempt < max_retries - 1:
                time.sleep(2)
    
    return None


def load_questions(xlsx_path: str) -> list:
    """Load questions from classified XLSX (includes category and image columns)."""
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    questions = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        questions.append({
            'codice': row[0].value,
            'domanda': row[1].value or '',
            'A': row[2].value or '',
            'B': row[3].value or '',
            'C': row[4].value or '',
            'D': row[5].value or '',
            'E': row[6].value or '',
            'categoria': row[7].value or '',
            'immagine': row[8].value or 'No',
            'cat_immagine': row[9].value or '',
        })
    return questions


# ==============================================================================
# Category Translation Mapping
# ==============================================================================

IT_TO_EN_CAT = {
    "Cardiologia e Cardiochirurgia": "Cardiology and Cardiac Surgery",
    "Chirurgia Generale": "General Surgery",
    "Gastroenterologia": "Gastroenterology",
    "Neurologia e Neurochirurgia": "Neurology and Neurosurgery",
    "Pneumologia e Chirurgia Toracica": "Pulmonology and Thoracic Surgery",
    "Ortopedia e Traumatologia": "Orthopedics and Traumatology",
    "Endocrinologia": "Endocrinology",
    "Ginecologia e Ostetricia": "Gynecology and Obstetrics",
    "Pediatria": "Pediatrics",
    "Urologia": "Urology",
    "Nefrologia": "Nephrology",
    "Ematologia": "Hematology",
    "Oncologia": "Oncology",
    "Malattie Infettive": "Infectious Diseases",
    "Immunologia e Reumatologia": "Immunology and Rheumatology",
    "Dermatologia e Venereologia": "Dermatology and Venereology",
    "Psichiatria": "Psychiatry",
    "Oftalmologia": "Ophthalmology",
    "Otorinolaringoiatria": "Otorhinolaryngology",
    "Anestesia e Rianimazione": "Anesthesia and Intensive Care",
    "Igiene, Epidemiologia e Statistica": "Hygiene, Epidemiology and Statistics",
    "Medicina del Lavoro": "Occupational Medicine",
    "Medicina Legale": "Forensic Medicine",
    "Diagnostica per Immagini e Medicina Nucleare": "Diagnostic Imaging and Nuclear Medicine",
    "Genetica Medica": "Medical Genetics",
    "Farmacologia e Tossicologia": "Pharmacology and Toxicology",
    "Nutrizione Clinica": "Clinical Nutrition",
    "Medicina Interna": "Internal Medicine",
}


def translate_category(cat_str: str) -> str:
    """Translate category string from Italian to English."""
    if not cat_str:
        return ""
    cats = [c.strip() for c in cat_str.split(";")]
    translated = [IT_TO_EN_CAT.get(c, c) for c in cats]
    return "; ".join(translated)


def process_year(year: int, client: anthropic.Anthropic) -> None:
    """
    Process all questions for a given year: load, translate, and save.

    Supports resuming from intermediate results if the process was interrupted.

    Args:
        year: Exam year (2017–2025).
        client: Anthropic API client instance.
    """
    xlsx_path = os.path.join(CLASSIFIED_DIR, f"{year}_answers_classified.xlsx")
    
    if not os.path.exists(xlsx_path):
        print(f"  File not found: {xlsx_path}")
        return
    
    questions = load_questions(xlsx_path)
    print(f"  Loaded {len(questions)} questions")
    
    # Check for intermediate results (resume support)
    intermediate_path = os.path.join(INTERMEDIATE_DIR, f"{year}_translations.json")
    existing_translations = {}
    if os.path.exists(intermediate_path):
        with open(intermediate_path, 'r', encoding='utf-8') as f:
            existing_translations = json.load(f)
        print(f"  Resuming from {len(existing_translations)} previously translated questions")
    
    all_translations = existing_translations.copy()
    
    for batch_start in range(0, len(questions), BATCH_SIZE):
        batch_end = min(batch_start + BATCH_SIZE, len(questions))
        
        # Check if batch already done (1-indexed)
        batch_keys = [str(i+1) for i in range(batch_start, batch_end)]
        if all(k in all_translations for k in batch_keys):
            continue
        
        batch = questions[batch_start:batch_end]
        print(f"  Translating questions {batch_start+1}-{batch_end}...")
        
        results = translate_batch(client, batch, batch_start)
        
        if results:
            for i, result in enumerate(results):
                q_idx = batch_start + i + 1  # 1-indexed
                all_translations[str(q_idx)] = {
                    'question': result.get('question', ''),
                    'A': result.get('A', ''),
                    'B': result.get('B', ''),
                    'C': result.get('C', ''),
                    'D': result.get('D', ''),
                    'E': result.get('E', ''),
                }
            
            # Save intermediate results
            with open(intermediate_path, 'w', encoding='utf-8') as f:
                json.dump(all_translations, f, ensure_ascii=False, indent=2)
        else:
            print(f"  FAILED to translate batch {batch_start+1}-{batch_end}")
        
        # Rate limiting
        time.sleep(1)
    
    # Create output xlsx
    output_path = os.path.join(OUTPUT_DIR, f"{year}_answers_english.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SSM_Test_EN"
    
    headers = ['Question Code', 'Question', 'Answer A', 'Answer B',
               'Answer C', 'Answer D', 'Answer E', 'Category',
               'Image', 'Image Category']
    ws.append(headers)
    
    for i, q in enumerate(questions):
        q_num = str(i + 1)
        translation = all_translations.get(q_num, {})
        
        # Use translated text if available, otherwise keep original
        question_en = translation.get('question', q['domanda'])
        a_en = translation.get('A', q['A'])
        b_en = translation.get('B', q['B'])
        c_en = translation.get('C', q['C'])
        d_en = translation.get('D', q['D'])
        e_en = translation.get('E', q['E'])
        
        # Translate category
        category_en = translate_category(q['categoria'])
        
        # Image column
        image_en = "Yes" if q['immagine'] == "Si" else "No"
        
        ws.append([
            q['codice'],
            question_en,
            a_en,
            b_en,
            c_en,
            d_en,
            e_en,
            category_en,
            image_en,
            q['cat_immagine'],  # Already in English
        ])
    
    wb.save(output_path)
    print(f"  Saved: {output_path}")
    
    # Quick stats
    translated_count = len(all_translations)
    print(f"  Translated: {translated_count}/{len(questions)}")


# ==============================================================================
# Main Execution
# ==============================================================================


def main():
    """
    Main entry point. Translates all SSM questions (2017–2025) from Italian
    to English using the Claude API.
    """
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print("ERROR: ANTHROPIC_API_KEY environment variable not set.")
        print("  Linux/macOS: export ANTHROPIC_API_KEY='sk-ant-...'")
        print("  Windows:     $env:ANTHROPIC_API_KEY = 'sk-ant-...'")
        sys.exit(1)

    print("=" * 70)
    print("ITAMed — Medical Question Translation (Italian → English)")
    print("=" * 70)
    print(f"  Model:            {MODEL}")
    print(f"  Input directory:  {CLASSIFIED_DIR}")
    print(f"  Output directory: {OUTPUT_DIR}")
    print(f"  Batch size:       {BATCH_SIZE}")
    print(f"  Years:            {YEARS}")
    print()

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(INTERMEDIATE_DIR, exist_ok=True)

    client = anthropic.Anthropic(api_key=api_key)

    for year in YEARS:
        print(f"\n--- {year} ---")
        process_year(year, client)

    print("\n" + "=" * 70)
    print("TRANSLATION COMPLETE")
    print(f"Output: {OUTPUT_DIR}")
    print("=" * 70)


if __name__ == "__main__":
    main()
